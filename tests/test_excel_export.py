from openpyxl import Workbook
from openpyxl import load_workbook

from excel_export import _format_data_sheet


def test_table_export_uses_table_filter_only():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Company ID", "Company Name"])
    worksheet.append([1, "Example Foundry GmbH"])

    _format_data_sheet(
        worksheet,
        table_name="CompaniesTable",
    )

    assert worksheet.auto_filter.ref is None
    assert worksheet.tables["CompaniesTable"].ref == "A1:B2"


def test_export_omits_logo_and_includes_normalized_categories(tmp_path, monkeypatch):
    import database
    from excel_export import export_exhibition_to_excel

    monkeypatch.setattr(database, "DATABASE_PATH", tmp_path / "test.db")
    monkeypatch.setattr(database, "DATA_DIR", tmp_path)
    database.initialize_database()
    exhibition_id = database.get_or_create_exhibition(
        "GIFA", 2023, "GMTN2023.gifa"
    )
    database.sync_exhibitor(
        {
            "exh": "EVENT.123",
            "name": "Example Foundry GmbH",
            "categories": [
                {
                    "id": "prod=event.01",
                    "label": "Casting",
                    "hierarchy": [
                        {"id": "01", "label": "Foundry equipment"}
                    ],
                }
            ],
        },
        exhibition_id,
    )

    output = tmp_path / "export.xlsx"
    export_exhibition_to_excel("GIFA", 2023, output)
    workbook = load_workbook(output)

    company_headers = [
        cell.value for cell in workbook["Company Master"][1]
    ]
    lead_headers = [
        cell.value for cell in workbook["Sales Leads"][1]
    ]
    assert "Logo" not in company_headers
    assert "Logo" not in lead_headers
    values = {
        header: workbook["Sales Leads"].cell(row=2, column=index + 1).value
        for index, header in enumerate(lead_headers)
    }
    assert values["Product Categories"] == "Casting"
    assert values["Product Groups"] == "Foundry equipment"